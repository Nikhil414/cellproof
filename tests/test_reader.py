import datetime
import time

import openpyxl
import pytest
from openpyxl.styles import Font
from xlsxkit import write_xlsx

from cellproof.model import Cell
from cellproof.reader import UnreadableWorkbook, load
from cellproof.compiler import NO_VALUE, compile_book
from cellproof.verifier import verify


def test_saved_empty_string_is_distinct_from_missing_answer(tmp_path):
    p = tmp_path / "empty.xlsx"
    write_xlsx(p, {"S": {
        "A1": ('=IF(TRUE,"","x")', ""),
        "A2": ('=A1&"suffix"', "suffix"),
        "B1": ("=1+1", None),
    }})
    book = load(p)
    assert (book.cells[("S", 1, 1)].value, book.cells[("S", 1, 1)].kind) == ("", "s")
    verdicts = {v.key: v for v in verify(book, compile_book(book))}
    assert verdicts[("S", 1, 1)].status == "PASS"
    assert verdicts[("S", 2, 1)].status == "PASS"
    assert verdicts[("S", 1, 2)].status == "UNRESOLVED"
    assert verdicts[("S", 1, 2)].reason == NO_VALUE


def test_reads_inputs_formulas_and_saved_values(tmp_path):
    p = tmp_path / "b.xlsx"
    write_xlsx(
        p,
        {"My Sheet": {
            "A1": 2, "B1": "hi", "C1": True,
            "A2": ("=A1*Rate", 4), "B2": ('=B1&"x"', "hix"), "C2": ("=1/0", "#DIV/0!"), "D2": ("=A1", None),
        }},
        names={"Rate": "'My Sheet'!$A$1"},
    )
    book = load(p)
    c = book.cells
    assert c[("My Sheet", 1, 1)] == Cell("My Sheet", 1, 1, None, 2.0, "n")
    assert (c[("My Sheet", 1, 2)].value, c[("My Sheet", 1, 2)].kind) == ("hi", "s")
    assert (c[("My Sheet", 1, 3)].value, c[("My Sheet", 1, 3)].kind) == (True, "b")
    assert c[("My Sheet", 2, 1)] == Cell("My Sheet", 2, 1, "A1*Rate", 4.0, "n")
    assert c[("My Sheet", 2, 2)] == Cell("My Sheet", 2, 2, 'B1&"x"', "hix", "s")
    assert c[("My Sheet", 2, 3)] == Cell("My Sheet", 2, 3, "1/0", "#DIV/0!", "e")
    assert c[("My Sheet", 2, 4)] == Cell("My Sheet", 2, 4, "A1", None, "blank")
    assert book.names == {"RATE": ("My Sheet", "A1")}
    assert book.sheets == ["My Sheet"] and book.max_row == {"My Sheet": 2}


def test_array_formula_is_flagged(tmp_path):
    p = tmp_path / "a.xlsx"
    write_xlsx(p, {"S": {"A1": 1, "B1": ("{=SUM(A1:A1*2)}", 2)}})
    cell = load(p).cells[("S", 1, 2)]
    assert cell.array and cell.formula == "SUM(A1:A1*2)"


def test_dates_become_excel_serials(tmp_path):
    p = tmp_path / "d.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active["A1"] = datetime.date(2026, 1, 1)
    wb.save(p)
    assert load(p).cells[("S", 1, 1)] == Cell("S", 1, 1, None, 46023.0, "n")


def test_rejects_unsupported_format(tmp_path):
    p = tmp_path / "old.xls"
    p.write_bytes(b"x")
    with pytest.raises(UnreadableWorkbook, match="save the workbook as .xlsx"):
        load(p)


def test_rejects_corrupt_file(tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_bytes(b"not a zip")
    with pytest.raises(UnreadableWorkbook, match="cannot open bad.xlsx"):
        load(p)


def test_rejects_a_file_that_unzips_too_large(tmp_path, monkeypatch):
    import cellproof.reader as reader
    p = tmp_path / "big.xlsx"
    write_xlsx(p, {"S": {"A1": "x" * 5000}})
    monkeypatch.setattr(reader, "MAX_UNZIPPED", 1000)
    with pytest.raises(UnreadableWorkbook, match="GB limit"):
        load(p)


def test_a_styled_cell_on_the_last_row_does_not_make_reading_slow(tmp_path):  # review: Focus 1 end to end
    p = tmp_path / "formatted.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 201):
        ws.cell(row=r, column=1, value=r)
    ws["G1048576"].font = Font(bold=True)  # an empty cell someone formatted at the bottom of the sheet
    wb.save(p)
    start = time.perf_counter()
    book = load(p)
    assert time.perf_counter() - start < 3  # walking every row x column took 6 s here, 24 s and 3.8 GB at column Z
    assert len(book.cells) == 200 and book.max_row["S"] == 1_048_576
