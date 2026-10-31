"""FILTER, SORT, UNIQUE spilled into a block, nested in each other (cellproof.fn.dynamic)."""
from xlsxkit import write_xlsx

from cellproof.compiler import compile_book
from cellproof.reader import load
from cellproof.verifier import verify

DATA = {"A1": "Rep", "B1": "Region", "C1": "Amount",
        "A2": "Asha", "B2": "West", "C2": 80.0,
        "A3": "Ravi", "B3": "East", "C3": 40.0,
        "A4": "Mina", "B4": "west", "C4": 60.0,
        "A5": "Vikram", "B5": "North", "C5": 40.0,
        "E1": 1.0, "F1": 0.0, "G1": 1.0, "H1": "a-b"}


def spill(anchor, formula, values, kind="dynamic"):
    """{anchor: ..., next cells: ...} for a block of `values` (a list of rows) starting at anchor."""
    from openpyxl.utils import column_index_from_string, get_column_letter
    from openpyxl.utils.cell import coordinate_from_string
    col, row = coordinate_from_string(anchor)
    c0 = column_index_from_string(col)
    last = f"{get_column_letter(c0 + len(values[0]) - 1)}{row + len(values) - 1}"
    out = {}
    for i, vals in enumerate(values):
        for j, v in enumerate(vals):
            ref = f"{get_column_letter(c0 + j)}{row + i}"
            out[ref] = ("{=" + formula + "}", v, f"{anchor}:{last}", kind) if (i, j) == (0, 0) else v
    return out


def check(tmp_path, cells):
    path = tmp_path / "d.xlsx"
    write_xlsx(path, {"S": {**DATA, **cells}})
    book = load(path)
    verdicts = {v.key: v for v in verify(book, compile_book(book))}
    return verdicts, {k: (v.status, v.expected, v.actual, v.reason) for k, v in verdicts.items()
                      if v.status != "PASS"}


def test_filter_rows_columns_and_empty(tmp_path):
    _, bad = check(tmp_path, {
        **spill("J1", '_xlfn._xlws.FILTER(A2:C5,B2:B5="West")', [["Asha", "West", 80.0], ["Mina", "west", 60.0]]),
        **spill("J5", "_xlfn._xlws.FILTER(A2:C5,C2:C5>50)", [["Asha", "West", 80.0], ["Mina", "west", 60.0]]),
        **spill("J8", '_xlfn._xlws.FILTER(A1:C1,E1:G1)', [["Rep", "Amount"]]),          # by columns
        **spill("J9", '_xlfn._xlws.FILTER(A2:A5,C2:C5>100,"none")', [["none"]]),
        **spill("J10", "_xlfn._xlws.FILTER(A2:A5,C2:C5>100,0)", [[0.0]]),
        **spill("J11", "_xlfn._xlws.FILTER(A2:A5,B2:B5)", [["#VALUE!"]]),               # text include
        **spill("J12", "_xlfn._xlws.FILTER(A2:A5,E1:G1)", [["#VALUE!"]]),               # include of another size
    })
    assert bad == {}


def test_sort_and_unique(tmp_path):
    _, bad = check(tmp_path, {
        **spill("L1", "_xlfn._xlws.SORT(C2:C5)", [[40.0], [40.0], [60.0], [80.0]]),
        **spill("M1", "_xlfn._xlws.SORT(A2:C5,3,-1)",  # equal keys keep their order (Ravi before Vikram)
                [["Asha", "West", 80.0], ["Mina", "west", 60.0], ["Ravi", "East", 40.0], ["Vikram", "North", 40.0]]),
        **spill("Q1", "_xlfn._xlws.SORT(A2:A5)", [["Asha"], ["Mina"], ["Ravi"], ["Vikram"]]),
        **spill("R1", "_xlfn.UNIQUE(B2:B5)", [["West"], ["East"], ["North"]]),           # text ignores case
        **spill("S1", "_xlfn.UNIQUE(C2:C5,,TRUE)", [[80.0], [60.0]]),                    # exactly once
        **spill("T1", "_xlfn._xlws.SORT(_xlfn.UNIQUE(B2:B5))", [["East"], ["North"], ["West"]]),
        **spill("U1", '_xlfn.UNIQUE(_xlfn._xlws.FILTER(C2:C5,C2:C5<70))', [[40.0], [60.0]]),
        **spill("V1", "_xlfn._xlws.SORT(C2:C5,1,2)", [["#VALUE!"]]),
    })
    assert bad == {}


def test_a_result_of_another_size_than_excels_block_fails(tmp_path):
    # a dynamic spill is exactly as big as Excel's result: one row short means the translation disagrees
    verdicts, bad = check(tmp_path, {**spill("J1", "_xlfn._xlws.SORT(C2:C5)", [[40.0], [40.0], [60.0]])})
    assert bad[("S", 1, 10)][:3] == ("FAIL", 40.0, "#SPILL!")
    # a legacy block bigger than the result shows #N/A past it, smaller ones keep the first rows
    _, bad = check(tmp_path, {**spill("J1", "_xlfn._xlws.SORT(C2:C3)", [[40.0], [80.0], ["#N/A"]], "legacy"),
                              **spill("K1", "_xlfn._xlws.SORT(C2:C5)", [[40.0], [40.0]], "legacy")})
    assert bad == {}


def test_what_excel_was_not_asked_is_not_guessed(tmp_path):
    verdicts, bad = check(tmp_path, {
        **spill("J1", "_xlfn._xlws.SORT(A2:B5,2)", [["Ravi", "East"], ["Vikram", "North"], ["Asha", "West"],
                                                    ["Mina", "west"]]),  # fine: letters only
        **spill("J6", "_xlfn._xlws.SORT(H1:H1)", [["a-b"]]),            # punctuation: Excel's own collation
        # nothing kept and no if_empty: #CALC! in Microsoft's docs, #VALUE! in an Excel-saved file
        **spill("K6", "_xlfn._xlws.FILTER(A2:A5,C2:C5>100)", [["#CALC!"]]),
        # empty sort keys, in a block bigger than the one element: UNRESOLVED, not a #SPILL! FAIL
        **spill("L6", "_xlfn._xlws.SORT(G1:G3)", [[0.0], [0.0], [1.0]]),
        "J9": ("=_xlfn._xlws.FILTER(A2:A5,C2:C5>50)", "Asha"),           # not an array formula
    })
    assert set(bad) == {("S", 6, 10), ("S", 6, 11), ("S", 6, 12), ("S", 7, 12), ("S", 8, 12), ("S", 9, 10)}
    assert all(v[0] == "UNRESOLVED" for v in bad.values())


def test_filter_sort_unique_inside_aggregates(tmp_path):
    _, bad = check(tmp_path, {
        "J1": ('=SUM(_xlfn._xlws.FILTER(C2:C5,B2:B5="West"))', 140.0),
        "J2": ("=COUNTA(_xlfn.UNIQUE(B2:B5))", 3.0),
        "J3": ("=ROWS(_xlfn.UNIQUE(C2:C5))", 3.0), "J4": ("=COLUMNS(A1:C5)", 3.0), "J5": ("=ROWS(A2:C5)", 4.0),
        "J6": ("=MAX(_xlfn._xlws.SORT(C2:C5))", 80.0),
        "J7": ('=AVERAGE(_xlfn._xlws.FILTER(C2:C5,C2:C5<70))', 140 / 3),
        "J8": ('=SUM(_xlfn._xlws.FILTER(C2:C5,C2:C5>100))', "#CALC!"),  # nothing kept, no if_empty
        "J9": ('=COUNT(_xlfn._xlws.FILTER(G1:G3,G1:G3<5))', 1.0),       # keeps two empty cells: not checked
    })
    assert set(bad) == {("S", 8, 10), ("S", 9, 10)} and all(v[0] == "UNRESOLVED" for v in bad.values())


def test_sortby_and_sequence(tmp_path):
    _, bad = check(tmp_path, {
        **spill("J1", "_xlfn.SORTBY(A2:A5,C2:C5,-1)", [["Asha"], ["Mina"], ["Ravi"], ["Vikram"]]),
        **spill("K1", "_xlfn.SORTBY(A2:A5,C2:C5,1,A2:A5,-1)", [["Vikram"], ["Ravi"], ["Mina"], ["Asha"]]),
        **spill("L1", "_xlfn.SORTBY(A2:A5,C2:C3)", [["#VALUE!"]]),                  # by_array of another height
        **spill("M1", "_xlfn.SEQUENCE(2,3)", [[1.0, 2.0, 3.0], [4.0, 5.0, 6.0]]),
        **spill("M4", "_xlfn.SEQUENCE(3,1,10,-5)", [[10.0], [5.0], [0.0]]),
        **spill("M8", "_xlfn.SEQUENCE(2.9)", [[1.0], [2.0]]),                        # truncated
        "M11": ("=SUM(_xlfn.SEQUENCE(100))", 5050.0),
        **spill("Q1", "_xlfn.SEQUENCE(0)", [["#CALC!"]]),                            # not checked
        **spill("Q2", "_xlfn.SEQUENCE(1E9,1E9)", [["#NUM!"]]),                       # never built
    })
    assert set(bad) == {("S", 1, 17), ("S", 2, 17)} and all(v[0] == "UNRESOLVED" for v in bad.values())


def test_corpus_findings_for_sortby_filter_and_xmatch(tmp_path):
    # formulas' test.xlsx (Excel-saved): SORTBY by one-row arrays sorts the columns; FILTER's if_empty range
    # spills whole; XMATCH next-smaller across types or over punctuation, and a sharp s, are not guessed
    _, bad = check(tmp_path, {
        **spill("J1", "_xlfn.SORTBY(C2:C3,{2},1)", [[80.0], [40.0]]),
        **spill("K1", "_xlfn.SORTBY(E1:G1,{3,1,2},1)", [[0.0, 1.0, 1.0]]),
        **spill("K3", "_xlfn.SORTBY({1,2,3;4,5,6},{3,2,2},1,{5,1,4},-1)", [[3.0, 2.0, 1.0], [6.0, 5.0, 4.0]]),
        **spill("K6", '_xlfn._xlws.FILTER(A2:A5,C2:C5>100,B2:B3)', [["West"], ["East"]]),
        "N1": ('=_xlfn.XMATCH("*b?u*",{-1.1,2.1,3.1,4.1},-1)', 4.0),
        "N2": ('=_xlfn.XMATCH("Mina",A2:A5,-1)', 3.0),             # same type, plain text: compiled
        "N3": ('=_xlfn.XMATCH("ss",{1,"ß"},0)', 2.0),
    })
    assert set(bad) == {("S", 1, 14), ("S", 3, 14)} and all(v[0] == "UNRESOLVED" for v in bad.values())
