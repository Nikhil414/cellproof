"""Load a workbook twice: once for formulas, once for the values Excel saved."""
from __future__ import annotations

import datetime
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
from openpyxl.worksheet.formula import ArrayFormula

from unexcel.model import Book, Cell, Kind, Value

ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
SUPPORTED = {".xlsx", ".xlsm"}
# An .xlsx is a zip of XML. A small file that unzips to gigabytes (a "zip bomb") would exhaust memory in openpyxl.
# ponytail: one total-size limit; raise it if real workbooks ever need more
MAX_UNZIPPED = 2 * 1024**3


class UnreadableWorkbook(Exception):
    pass


def _saved(cell, epoch) -> tuple[Value, Kind]:
    v = cell.value
    if v is None:
        if cell.data_type == "str":  # Excel's cached empty string is None in openpyxl.
            return "", "s"
        return None, "blank"
    if cell.data_type == "e":
        return str(v), "e"
    if isinstance(v, bool):  # before int: bool is an int subclass
        return v, "b"
    if isinstance(v, (int, float)):
        return float(v), "n"
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return float(to_excel(v, epoch)), "n"  # handles Excel's 1900 leap-year bug and the 1904 system
    return str(v), "s"


def load(path: str | Path) -> Book:
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED:
        raise UnreadableWorkbook(f"{path.suffix or 'these'} files are not supported in v1 - save the workbook as .xlsx")
    try:
        with zipfile.ZipFile(path) as z:
            size = sum(i.file_size for i in z.infolist())
    except (OSError, zipfile.BadZipFile) as err:
        raise UnreadableWorkbook(f"cannot open {path.name}: {err}") from err
    if size > MAX_UNZIPPED:
        raise UnreadableWorkbook(f"cannot open {path.name}: it unzips to {size / 1024**3:.1f} GB, more than the "
                                 f"{MAX_UNZIPPED // 1024**3} GB limit")
    try:
        fwb = openpyxl.load_workbook(path)
        vwb = openpyxl.load_workbook(path, data_only=True)
    except Exception as err:  # openpyxl raises many types for corrupt or encrypted files
        raise UnreadableWorkbook(f"cannot open {path.name}: {err}") from err

    names = {}
    for name, dn in fwb.defined_names.items():
        dests = list(dn.destinations)
        if len(dests) == 1:
            sheet, ref = dests[0]
            names[name.upper()] = (sheet, ref.replace("$", ""))

    cells, max_row = {}, {}
    for ws in fwb.worksheets:
        vws = vwb[ws.title]
        max_row[ws.title] = ws.max_row
        # ponytail: openpyxl's private _cells holds only the cells the file has; iter_rows() would build
        # every row x column up to max_row (one styled cell at G1048576 cost 6 s; at Z, 3.8 GB)
        for (r, col), c in sorted(ws._cells.items()):
            if c.value is None:
                continue
            value, kind = _saved(vws.cell(row=r, column=col), fwb.epoch)
            formula, array = None, False
            if isinstance(c.value, ArrayFormula):
                formula, array = (c.value.text or "").removeprefix("="), True
            elif c.data_type == "f":
                formula = c.value.removeprefix("=")
            cells[(ws.title, c.row, c.column)] = Cell(ws.title, c.row, c.column, formula, value, kind, array)
    return Book(cells, [ws.title for ws in fwb.worksheets], names, max_row, fwb.epoch == CALENDAR_MAC_1904)
