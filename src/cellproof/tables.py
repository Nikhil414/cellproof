"""Structured references to Excel tables (Sales[Amount], Sales[[#This Row],[Qty]], Sales[#Totals]) -> A1 text.

The reader rewrites them before anything else reads a formula, so dependencies, regions and the compiler only ever
see plain references: a column is an absolute range ('S'!$C$2:$C$9), and [#This Row] (written [@Qty] in Excel) is
the cell in the formula's own row with a relative row ('S'!$C5), so a table column pasted down is a region.
Files store the long form (Table1[[#This Row],[Qty]]); the short forms are read too.
"""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import get_column_letter

from cellproof.model import Unsupported

SPECIALS = {"#ALL", "#DATA", "#HEADERS", "#TOTALS", "#THIS ROW"}


@dataclass(frozen=True)
class Table:
    name: str
    sheet: str
    r1: int
    c1: int
    r2: int
    c2: int
    header: int  # header rows (0 or 1)
    totals: int  # totals rows (0 or 1)
    columns: tuple[str, ...]

    def contains(self, sheet: str, row: int, col: int) -> bool:
        return sheet == self.sheet and self.r1 <= row <= self.r2 and self.c1 <= col <= self.c2


def _bracket(text: str, i: int) -> tuple[str, int]:
    """The name in the bracket at text[i] == '[' (with ' escapes undone) and the index after its ']'."""
    out, i = [], i + 1
    while i < len(text):
        ch = text[i]
        if ch == "'" and i + 1 < len(text):
            out.append(text[i + 1])
            i += 2
        elif ch == "]":
            return "".join(out), i + 1
        elif ch == "[":
            raise Unsupported(f"structured reference {text} not understood")
        else:
            out.append(ch)
            i += 1
    raise Unsupported(f"structured reference {text} not understood")


def _items(inner: str, text: str) -> list[tuple[str, ...]]:
    """[a], [b]:[c] ... inside the outer brackets -> [("a",), ("b", "c")]; plain text -> [(text,)]."""
    s = inner.strip()
    if s.startswith("@"):  # Excel 2010 short form: [@Qty], [@[Unit Price]], [@]
        rest = s[1:].strip()
        return [("#This Row",)] + (_items(rest, text) if rest.startswith("[") else [(rest,)] if rest else [])
    if not s.startswith("["):
        return [(s.replace("'", ""),)] if s else []
    items, i = [], 0
    while i < len(s):
        if s[i] in " ,":
            i += 1
            continue
        if s[i] != "[":
            raise Unsupported(f"structured reference {text} not understood")
        a, i = _bracket(s, i)
        j = i
        while j < len(s) and s[j] == " ":
            j += 1
        if j < len(s) and s[j] == ":":
            j += 1
            while j < len(s) and s[j] == " ":
                j += 1
            b, i = _bracket(s, j)
            items.append((a, b))
        else:
            items.append((a,))
    return items


def _rows(t: Table, specials: set[str], row: int, text: str) -> tuple[int, int] | None:
    """The table rows a reference covers; None for [#This Row] outside the data rows (Excel: #VALUE!)."""
    first, last = t.r1 + t.header, t.r2 - t.totals
    if specials == {"#THIS ROW"}:
        return (row, row) if first <= row <= last else None
    if "#THIS ROW" in specials:
        raise Unsupported(f"structured reference {text} mixes [#This Row] with other rows")
    if "#HEADERS" in specials and not t.header or "#TOTALS" in specials and not t.totals:
        raise Unsupported(f"structured reference {text}: table {t.name} has no such row")
    if not specials or specials == {"#DATA"}:
        return first, last
    if specials == {"#ALL"}:
        return t.r1, t.r2
    if specials == {"#HEADERS"}:
        return t.r1, t.r1
    if specials == {"#TOTALS"}:
        return t.r2, t.r2
    if specials == {"#HEADERS", "#DATA"}:
        return t.r1, last
    if specials == {"#DATA", "#TOTALS"}:
        return first, t.r2
    raise Unsupported(f"structured reference {text} not understood")


def _column(t: Table, name: str, text: str) -> int:
    for i, col in enumerate(t.columns):
        if col.strip().lower() == name.strip().lower():
            return t.c1 + i
    raise Unsupported(f"structured reference {text}: table {t.name} has no column {name!r}")


def to_a1(text: str, tables: dict[str, Table], at: tuple[str, int, int]) -> str:
    """One structured reference token -> A1 text (or #VALUE! for [#This Row] outside the table's data rows)."""
    k = text.index("[")
    name = text[:k].strip()
    if name:
        t = tables.get(name.upper())
        if t is None:
            raise Unsupported(f"unknown table {name!r}")
    else:  # [Qty] written inside its own table
        t = next((t for t in tables.values() if t.contains(*at)), None)
        if t is None:
            raise Unsupported(f"structured reference {text} outside any table")
    inner, end = text[k + 1:-1], text.rstrip()
    if not end.endswith("]"):
        raise Unsupported(f"structured reference {text} not understood")
    items = _items(inner, text)
    specials = {i[0].upper() for i in items if len(i) == 1 and i[0].upper() in SPECIALS}
    columns = [i for i in items if not (len(i) == 1 and i[0].upper() in SPECIALS)]
    if len(columns) > 1:
        raise Unsupported(f"structured reference {text} with several column items not supported yet")
    if columns:
        ends = [_column(t, n, text) for n in columns[0]]
        c1, c2 = min(ends), max(ends)
    else:
        c1, c2 = t.c1, t.c2
    rows = _rows(t, specials, at[1], text)
    if rows is None:
        return "#VALUE!"
    sheet = "" if t.sheet == at[0] else "'" + t.sheet.replace("'", "''") + "'!"  # own sheet: as Excel writes it
    if specials == {"#THIS ROW"}:
        cells = [f"${get_column_letter(c)}{rows[0]}" for c in (c1, c2)]  # relative row: pasted down, same shape
    else:
        cells = [f"${get_column_letter(c1)}${rows[0]}", f"${get_column_letter(c2)}${rows[1]}"]
    return sheet + (cells[0] if cells[0] == cells[1] else f"{cells[0]}:{cells[1]}")


def expand(formula: str, tables: dict[str, Table], at: tuple[str, int, int]) -> str:
    """The formula (without '=') with every structured reference rewritten as A1 text."""
    if "[" not in formula:
        return formula
    out, changed = [], False
    for t in Tokenizer("=" + formula).items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE and "[" in t.value and "!" not in t.value:
            out.append(to_a1(t.value, tables, at))
            changed = True
        else:
            out.append(t.value)
    return "".join(out) if changed else formula  # the tokens drop some line breaks: rebuild only when needed
