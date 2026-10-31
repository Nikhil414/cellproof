"""Formula AST -> SQL. Cell mode, run in batches.

Every cell known so far (inputs, then the results of earlier batches) lives in one table with the
columns of INPUT_COLUMNS: kind is n/s/b/e, num holds numbers and TRUE/FALSE as 1/0, txt holds text
and Excel error codes. A batch is up to BATCH consecutive formula cells (or BATCH_WEIGHT characters of SQL) in dependency order. It
becomes one SQL statement: the cells it computes are materialized CTEs (v = value, e = error code
or NULL), appended to the table it read from. One statement per workbook broke DuckDB at ~1k
formula cells; batches keep every statement small.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from dataclasses import field as dc_field
from pathlib import Path

import cellproof.fn  # noqa: F401  (registers the function library)
from openpyxl.utils import range_boundaries

from cellproof.expr import (BLANK, CELLS, NO_ERR, Expr, any_of, compare, error, errs, finalize, lit_num, lit_str, num,
                          txt, where_cell)
from cellproof.functions import FUNCS
from cellproof.graph import formula_order
from cellproof.model import Book, Cell, Key, Unsupported, a1
from cellproof.parser import ARRAYS, Binary, Bool, Call, ErrLit, Num, Postfix, RefNode, Str, Unary, arrays, parse
from cellproof.rangeindex import RangeIndex, field
from cellproof.refs import Area, resolve

INPUT_COLUMNS = "sheet VARCHAR, row_num INTEGER, col_num INTEGER, kind VARCHAR, num DOUBLE, txt VARCHAR"
SEED_HEADER = ["sheet", "row_num", "col_num", "kind", "num", "txt"]
READ_SEED = ("read_csv(?, header = true, quote = '\"', escape = '\"', columns = {'sheet': 'VARCHAR', "
             "'row_num': 'INTEGER', 'col_num': 'INTEGER', 'kind': 'VARCHAR', 'num': 'DOUBLE', 'txt': 'VARCHAR'})")
VOLATILE = {"NOW", "TODAY", "RAND", "RANDBETWEEN", "OFFSET", "INDIRECT"}
PASS_THROUGH = {"IF", "IFS", "IFERROR", "IFNA", "CHOOSE", "SWITCH"}  # can return a range given as an argument
ARITH = {"+", "-", "*"}
KIND = {"num": "n", "txt": "s", "bool": "b", "err": "e", "any": "a"}  # a: decided when the SQL runs
TYPE = {"n": "num", "s": "txt", "b": "bool"}
NO_VALUE = "no saved value in the file - open and save it in Excel"
DEEP = "formula is nested too deeply to translate"
TOO_BIG = "formula's SQL would be too large to run"
# IFERROR/division repeat their arguments in the value and the error SQL, so SQL can double per
# nesting level; checking every node keeps memory near this bound (200k-cell workbooks stay < 10 KB/cell)
MAX_SQL = 1_000_000
# A batch closes at BATCH cells or when its SQL weighs BATCH_WEIGHT characters, whichever comes first. Time inside one
# statement grows faster than linearly with its cell count, and lookups are far heavier than plain arithmetic, so a
# count alone makes lookup-heavy batches slow (measured on DuckDB 1.5.6: 25 VLOOKUP cells 0.2 s, 250 cells 17.8 s).
# BATCH also keeps the UNION ALL chain below DuckDB's expression depth limit of 1000.
BATCH = 250
# Sweep (tests/bench_shapes.py, verify seconds / batches; vlookup 500, running_input 2000, flat 5000; sum):
#   W=5000:  3.3/167  3.2/167  7.7/358  sum 14.2
#   W=10000: 3.0/84   3.1/84   7.2/186  sum 13.3  <- chosen (smallest sum)
#   W=20000: 3.2/46   3.9/43   8.3/93   sum 15.4
#   W=40000: 4.0/24   4.8/22   12.0/47  sum 20.8
#   W=80000: 5.2/12   10.1/11  25.9/24  sum 41.2
BATCH_WEIGHT = 10_000  # characters
ROW_COST = 250  # characters a cell adds besides its expression: its CTE wrapper and the row it emits


@dataclass(frozen=True)
class Step:
    key: Key
    expr: Expr
    cte: str  # its CTE name inside its batch


def step_weight(e: Expr) -> int:
    return ROW_COST + len(e.sql) + len(e.err) + len(e.txt) + len(e.kind)


@dataclass
class Compiled:
    batches: list[list[Step]]
    unresolved: dict[Key, str]
    findings: list[str]
    notes: list[str] = dc_field(default_factory=list)
    chains: dict[Key, Key] = dc_field(default_factory=dict)  # cell -> the unresolved cell that stopped its compile

    @property
    def steps(self) -> list[Step]:
        return [s for batch in self.batches for s in batch]


class NewBatch(Exception):
    """A range reads a cell of the open batch. Ranges read the table, so the batch must close first."""


class Keys:
    """The non-empty cells of a range, summarised in O(log n). Iterating lists them (O(range): avoid on big ranges)."""

    def __init__(self, comp: "Compiler", area: Area, present: int, total: int):
        self._comp, self.area, self._n = comp, area, present
        self.kinds = frozenset(name for name in ("n", "s", "b", "e", "a", "blank") if field(total, name))
        self.may_err = field(total, "merr") > 0

    def __len__(self) -> int:
        return self._n

    def __iter__(self):
        return iter(self.area.present(self._comp.book))

    def text_cells(self) -> list:
        """Input text cells in the range, row-major (the only cells SUM-like functions must warn about)."""
        a = self.area
        return self._comp.index.text_cells(a.sheet, a.r1, a.c1, a.r2, a.c2)


# Analysis ToolPak functions do not take an array where one value is expected: {arguments that are one value}
NO_LIFT = {"EDATE": 2, "EOMONTH": 2, "MROUND": 2, "YEARFRAC": 3, "NETWORKDAYS": 2, "WORKDAY": 2}
NO_LIFT_CHECKED = {"EDATE", "EOMONTH", "MROUND", "YEARFRAC"}  # #VALUE! there, in formulas' Excel-saved test.xlsx


class Compiler:
    def __init__(self, book: Book):
        add_arrays(book)
        self.book = book
        self.done: dict[Key, tuple[str, bool]] = {}  # formula cell in a closed batch -> (type, may be an error)
        self.open: dict[Key, tuple[str, str, bool]] = {}  # formula cell in the open batch -> (cte, type, may err)
        self.findings: list[str] = []
        self.index = RangeIndex(book)
        self.rowrefs: dict[str, Expr] | None = None  # same-row placeholders of a region (see regions)
        self.at: Key | None = None  # the formula cell being compiled in cell mode (implicit intersection)
        self.ctx: list[str] = []  # what consumes the node being compiled: "op", "fn" or "pass" (PASS_THROUGH)
        self.wid = 0  # unique aliases w1, w2... for the element sources of one array argument
        self.spill: tuple[int, int] | None = None  # array formula: this cell's (row, col) offset in its block
        self.spill_dims: list[tuple[int, int]] = []  # sizes of the ranges taken element by element (see _element)

    def walias(self) -> str:
        self.wid += 1
        return f"w{self.wid}"

    def compile_cell(self, cell: Cell) -> Expr:
        if cell.key in self.book.unread:
            raise Unsupported(self.book.unread[cell.key])
        self.at = cell.key
        self.spill, self.spill_dims = None, []
        if cell.array:
            c1, r1, _, _ = range_boundaries(cell.ref or a1(cell.key).split("!")[-1])
            self.spill = (cell.row - r1, cell.col - c1)
        try:
            e = finalize(self.expr(parse(cell.formula), cell.sheet))
            if self.spill is not None and self.spill_dims:
                # real Excel (formulas' test.xlsx): the array is as big as its biggest range, one row (column)
                # repeating; a block cell beyond it is #N/A whatever the formula does with errors -
                # {=IFERROR(G9:Q9,R9)} one cell past Q9 is #N/A, not R9
                rows = max(d[0] for d in self.spill_dims)
                cols = max(d[1] for d in self.spill_dims)
                i, j = self.spill
                if (rows > 1 and i >= rows) or (cols > 1 and j >= cols):
                    return error("#N/A")
            return e
        finally:
            self.at, self.spill = None, None

    def _intersect(self, area, text: str) -> Expr:
        """A range where one value is expected: Excel's implicit intersection takes the cell in the formula's own row
        (a one-column range) or column (a one-row range), and is #VALUE! when there is none."""
        if self.spill is not None:
            return self._element(area)
        if area.sheet in self.book.arrays:
            raise Unsupported("array constant used where one value is expected (needs array evaluation) not supported yet")
        if self.at is None or area.sheet != self.at[0] or (area.r1 != area.r2 and area.c1 != area.c2):
            raise Unsupported(f"range {text} used as a single value")
        # real Excel: in MAX(IF(c, A1:A9, 0)) IF hands MAX the whole range; only an operator, a function's own
        # argument or the cell itself takes one value of it
        for consumer in reversed(self.ctx):
            if consumer == "op":
                break
            if consumer == "fn":
                if self.ctx[-1] == "pass":
                    raise Unsupported(f"range {text} passed through IF/CHOOSE into another function (Excel may use "
                                      "the whole range there) not supported yet")
                break
        _, row, col = self.at
        if area.c1 == area.c2 and area.r1 <= row <= area.r2:
            return self.cell_value((area.sheet, row, area.c1))
        if area.r1 == area.r2 and area.c1 <= col <= area.c2:
            return self.cell_value((area.sheet, area.r1, col))
        return error("#VALUE!")

    def element_key(self, area) -> Key | None:
        """In an array formula a range where one value is expected is taken element by element: this cell's element,
        the only row (column) of a range one row (column) high for every row (column); None beyond its size, where
        Excel pads the range with #N/A (so ISERROR there is TRUE: real Excel, formulas' test.xlsx)."""
        rows, cols = area.r2 - area.r1 + 1, area.c2 - area.c1 + 1
        i, j = self.spill
        i, j = (0 if rows == 1 else i), (0 if cols == 1 else j)
        return None if i >= rows or j >= cols else (area.sheet, area.r1 + i, area.c1 + j)

    def _element(self, area, sized: bool = True) -> Expr:
        """sized=False: an IS function's own argument, which real Excel pads with #N/A without sizing the array
        ({=ISERROR(C8:Q8)} past Q8 is TRUE)."""
        if sized:
            self.spill_dims.append((area.r2 - area.r1 + 1, area.c2 - area.c1 + 1))
        key = self.element_key(area)
        return error("#N/A") if key is None else self.cell_value(key)

    def expr(self, node, sheet: str) -> Expr:
        e = self._expr(node, sheet)
        if sum(len(part) for part in (e.sql, e.err, e.txt, e.kind)) > MAX_SQL:
            raise Unsupported(TOO_BIG)
        return e

    def _expr(self, node, sheet: str) -> Expr:
        if node is None:
            return BLANK
        if isinstance(node, Num):
            return lit_num(node.value)
        if isinstance(node, Str):
            return lit_str(node.value)
        if isinstance(node, Bool):
            return Expr("TRUE" if node.value else "FALSE", "bool")
        if isinstance(node, ErrLit):
            return error(node.value)
        if isinstance(node, RefNode):
            if self.rowrefs is not None and node.text in self.rowrefs:
                return self.rowrefs[node.text]
            area = resolve(node.text, sheet, self.book)
            if not area.is_single:
                return self._intersect(area, node.text)
            return self.cell_value(area.key)
        kind = "op" if isinstance(node, (Unary, Postfix, Binary)) else "pass" if (
            isinstance(node, Call) and node.name in PASS_THROUGH) else "fn"
        self.ctx.append(kind)
        try:
            return self._compound(node, sheet)
        finally:
            self.ctx.pop()

    def _compound(self, node, sheet: str) -> Expr:
        if isinstance(node, Unary):
            x = num(self.expr(node.operand, sheet))
            return Expr(f"(-{x.sql})", "num", x.err)
        if isinstance(node, Postfix):
            x = num(self.expr(node.operand, sheet))
            return Expr(f"({x.sql} / 100)", "num", x.err)
        if isinstance(node, Binary):
            return self.binary(node, sheet)
        if isinstance(node, Call):
            if node.name in VOLATILE:
                raise Unsupported(f"volatile function {node.name}")
            fn = FUNCS.get(node.name)
            if fn is None:
                raise Unsupported(f"function {node.name} not supported yet")
            if self.spill is not None and node.name in NO_LIFT:
                ranged = [i for i, a in enumerate(node.args) if i < NO_LIFT[node.name] and isinstance(a, RefNode)
                          and not resolve(a.text, sheet, self.book).is_single]
                if ranged and node.name in NO_LIFT_CHECKED:
                    return error("#VALUE!")  # real Excel: {=EDATE(E9:I9,C9:G9)} is #VALUE! in every cell
                if ranged:
                    raise Unsupported(f"{node.name} of a range in an array formula not supported yet")
            return fn(self, node.args, sheet)
        raise Unsupported(f"cannot translate {node!r}")

    def binary(self, node: Binary, sheet: str) -> Expr:
        a, b, op = self.expr(node.left, sheet), self.expr(node.right, sheet), node.op
        if op in ARITH or op in ("/", "^"):
            x, y = num(a), num(b)
            if op in ARITH:
                return Expr(f"({x.sql} {op} {y.sql})", "num", errs(x.err, y.err))
            if op == "/":  # DuckDB gives inf, Excel #DIV/0!
                return Expr(f"({x.sql} / NULLIF({y.sql}, 0))", "num",
                            errs(x.err, y.err, f"(CASE WHEN {y.sql} = 0 THEN '#DIV/0!' END)"))
            # DuckDB POWER returns inf or nan where Excel returns an error (verified on 1.5.6)
            # Real Excel: (-8)^(1/3) is -1.9999999999999998, the real cube root: a negative number to the power of
            # 1/odd is minus the root of its absolute value. Any other fractional power of a negative number is #NUM!.
            root = f"(1 / NULLIF({y.sql}, 0))"
            odd_root = (f"({x.sql} < 0 AND {y.sql} <> FLOOR({y.sql}) AND abs({root} - round({root})) < 1e-9 "
                        f"AND abs(round({root})) % 2 = 1)")
            bad = (f"(CASE WHEN {x.sql} = 0 AND {y.sql} < 0 THEN '#DIV/0!' "
                   f"WHEN {x.sql} = 0 AND {y.sql} = 0 THEN '#NUM!' "
                   f"WHEN {x.sql} < 0 AND {y.sql} <> FLOOR({y.sql}) AND NOT {odd_root} THEN '#NUM!' END)")
            return Expr(f"(CASE WHEN {odd_root} THEN -POWER(-{x.sql}, {y.sql}) ELSE POWER({x.sql}, {y.sql}) END)",
                        "num", errs(x.err, y.err, bad))
        if op == "&":
            x, y = txt(a), txt(b)
            return Expr(f"({x.sql} || {y.sql})", "txt", errs(x.err, y.err))
        return compare(a, b, op)

    def cell_value(self, key: Key) -> Expr:
        cell = self.book.cells.get(key)
        if cell is None or (cell.formula is None and cell.kind == "blank"):
            return BLANK
        if key in self.open:
            cte, typ, may_err = self.open[key]
            err = f"(SELECT e FROM {cte})" if may_err else NO_ERR
            if typ == "any":
                return any_of(f"(SELECT k FROM {cte})", f"(SELECT v FROM {cte})", f"(SELECT t FROM {cte})", err)
            return Expr(f"(SELECT v FROM {cte})", typ, err)
        if cell.formula is None:
            if cell.kind == "e":
                return error(cell.value)
            typ, may_err = TYPE[cell.kind], False
        elif key in self.done:
            typ, may_err = self.done[key]
        else:
            raise Unsupported(f"depends on unresolved cell {a1(key)}", key)
        where = f"FROM {CELLS} WHERE {where_cell(key)}"
        err = f"(SELECT CASE WHEN kind = 'e' THEN txt END {where})" if may_err else NO_ERR
        if typ == "num":
            return Expr(f"(SELECT num {where})", "num", err)
        if typ == "txt":
            return Expr(f"(SELECT COALESCE(txt, '') {where})", "txt", err)  # the seed CSV reads '' back as NULL
        if typ == "bool":
            return Expr(f"((SELECT num {where}) <> 0)", "bool", err)
        if typ == "any":
            return any_of(f"(SELECT kind {where})", f"(SELECT num {where})", f"(SELECT COALESCE(txt, '') {where})", err)
        return Expr("NULL", "err", err)

    def ready(self, area: Area) -> Keys:
        """The non-empty cells of a range that SQL will read from the table. All must be computed already."""
        s = self.index.summary(area.sheet, area.r1, area.c1, area.r2, area.c2)
        if field(s.total, "fdone") < s.formulas:
            self._blocked(area)
        return Keys(self, area, s.present, s.total)

    def _blocked(self, area: Area) -> None:
        """Some formula cell of the range is unfinished: raise for the first one, in row-major order."""
        for key in area.present(self.book):
            if self.book.cells[key].formula is None:
                continue
            if key in self.open:
                raise NewBatch
            if key not in self.done:
                raise Unsupported(f"depends on unresolved cell {a1(key)}", key)
        raise AssertionError("range index is out of step with the compiler")

    def kind_of(self, key: Key) -> str:
        """Value kind (n/s/b/e) a cell has in the table when it isn't an error. Call ready() first."""
        cell = self.book.cells[key]
        return KIND[self.done[key][0]] if cell.formula is not None else cell.kind

    def may_err(self, key: Key) -> bool:
        """Whether the table can hold an error for this cell. Call ready() first."""
        cell = self.book.cells[key]
        return self.done[key][1] if cell.formula is not None else cell.kind == "e"


def add_arrays(book: Book) -> None:
    """The values of the book's array constants as cells of their hidden sheets (see parser.ARRAY_SHEET)."""
    for cell in list(book.cells.values()):
        if cell.formula is None or "{" not in cell.formula:
            continue
        try:
            names = set(arrays(parse(cell.formula)))
        except (Unsupported, RecursionError):
            continue
        for name in names - book.arrays:
            book.arrays.add(name)
            for r, row in enumerate(ARRAYS[name], 1):
                for c, (kind, value) in enumerate(row, 1):
                    book.cells[(name, r, c)] = Cell(name, r, c, None, value, kind)
            book.max_row[name] = len(ARRAYS[name])
            book.__dict__.pop("_rows_by_col", None)  # cached from the cells


def _finish(comp: Compiler, batches: list, batch) -> None:
    """A batch is finished: later cells read its results from the table. `batch` is a list of Steps or a region batch."""
    batches.append(batch)
    comp.done.update((s.key, (s.expr.type, s.expr.err != NO_ERR)) for s in batch)
    for s in batch:
        comp.index.complete(s.key, KIND[s.expr.type], s.expr.err != NO_ERR)


def _close(comp: Compiler, batches: list[list[Step]], current: list[Step]) -> None:
    """Seal the open batch: later cells read its results from the table."""
    if current:
        _finish(comp, batches, list(current))
        comp.open.clear()
        current.clear()


def compile_book(book: Book, regions: bool = True) -> Compiled:
    from openpyxl.utils import get_column_letter

    from cellproof.region import compile_region, region_note
    from cellproof.regions import RegionScan, column_order, find_regions, order_with_regions

    add_arrays(book)
    scan = find_regions(book) if regions else RegionScan([], [])
    order, unresolved, kept = order_with_regions(book, scan.regions)
    comp = Compiler(book)
    chains: dict[Key, Key] = {}  # cell -> the unresolved cell that stopped its compile (see impact.py)
    if book.app and not (book.app.startswith("Microsoft") and "Compatible" not in book.app):  # openpyxl: "Compatible"
        comp.findings.append(f"the file was last saved by {book.app}, not Microsoft Excel: its saved values follow that "
                             "program's rules, so a FAIL can be a difference between that program and Excel")
    batches: list = []
    current: list[Step] = []
    weight = 0
    first = {r.first_key: r for r in kept}
    inside = {k for r in kept for k in r.keys()} - set(first)
    notes = list(scan.rejected)
    kept_sets = [set(r.keys()) for r in kept]
    for r in scan.regions:
        if set(r.keys()) in kept_sets:
            continue
        letters = lambda cols: ", ".join(get_column_letter(c) for c in cols)  # noqa: E731
        pieces = [k for k in kept if (k.sheet, k.r1, k.r2) == (r.sheet, r.r1, r.r2) and set(k.cols) <= set(r.cols)]
        lost = sorted(set(r.cols) - {c for k in pieces for c in k.cols})
        why = "another cell sits between them in the dependency order"
        if pieces:
            blocks = (f"split into blocks ({' | '.join(letters(k.cols) for k in pieces)})" if len(pieces) > 1
                      else f"kept columns {letters(pieces[0].cols)} as a block")
            notes.append(f"{r.sheet} rows {r.r1}-{r.r2} {blocks}: {why}"
                         + (f"; columns {letters(lost)} stay in cell mode" if lost else ""))
        else:
            notes.append(f"{r.sheet} rows {r.r1}-{r.r2}, columns {letters(r.cols)} stay in cell mode: they could not "
                         f"be ordered as one block ({why})")
    used_models: set[str] = set()

    def add(key) -> None:
        nonlocal weight
        cell = book.cells[key]
        if cell.kind == "blank":
            unresolved[key] = NO_VALUE
            return
        if current and (len(current) >= BATCH or weight >= BATCH_WEIGHT):
            _close(comp, batches, current)
            weight = 0
        try:
            try:
                e = comp.compile_cell(cell)
            except NewBatch:
                _close(comp, batches, current)
                weight = 0
                e = comp.compile_cell(cell)  # the open batch is empty now, so this can't raise NewBatch
        except Unsupported as err:
            unresolved[key] = str(err)
            if err.dep is not None:
                chains[key] = err.dep
            return
        except RecursionError:  # Python's stack, not DuckDB's: DuckDB took 900 nesting levels in tests
            unresolved[key] = DEEP
            return
        step = Step(key, e, f"c{len(current) + 1}")
        comp.open[key] = (step.cte, e.type, e.err != NO_ERR)
        current.append(step)
        weight += step_weight(e)

    for key in order:
        if key in inside:
            continue
        region = first.get(key)
        if region is None:
            add(key)
            continue
        _close(comp, batches, current)  # the region reads finished cells from the table
        weight = 0
        batch, rejected = compile_region(comp, region)
        if batch is not None:
            _finish(comp, batches, batch)
        if batch is not None:
            base, i = batch.model, 1
            while batch.model in used_models:
                i += 1
                batch.model = f"{base}_{i}"
            used_models.add(batch.model)
        notes += region_note(region, batch, rejected)
        for col in column_order(region.cols, region.reads)[0]:
            if col in rejected:
                for row in range(region.r1, region.r2 + 1):
                    add((region.sheet, row, col))
    _close(comp, batches, current)
    return Compiled(batches, unresolved, comp.findings, notes, chains)


def cell_row(typ: str, v: str, e: str, t: str, k: str) -> tuple[str, str, str]:
    """The kind, num and txt SQL of one result row, from SQL for its value v, error e, text t and kind k."""
    clean = e == NO_ERR  # the cell cannot hold an error: no CASE needed
    if typ == "any":  # the kind is only known when the SQL runs
        if clean:
            return k, f"CASE WHEN {k} <> 's' THEN CAST({v} AS DOUBLE) END", f"CASE WHEN {k} = 's' THEN {t} END"
        return (f"CASE WHEN {e} IS NULL THEN {k} ELSE 'e' END",
                f"CASE WHEN {e} IS NULL AND {k} <> 's' THEN CAST({v} AS DOUBLE) END",
                f"CASE WHEN {e} IS NOT NULL THEN {e} WHEN {k} = 's' THEN {t} END")
    if clean:
        num_sql = "CAST(NULL AS DOUBLE)" if typ in ("txt", "err") else f"CAST({v} AS DOUBLE)"
        return f"'{KIND[typ]}'", num_sql, (f"CAST({v} AS VARCHAR)" if typ == "txt" else "NULL")
    num_sql = ("CAST(NULL AS DOUBLE)" if typ in ("txt", "err")
               else f"CASE WHEN {e} IS NULL THEN CAST({v} AS DOUBLE) END")
    txt_sql = f"COALESCE({e}, CAST({v} AS VARCHAR))" if typ == "txt" else e
    return f"CASE WHEN {e} IS NULL THEN '{KIND[typ]}' ELSE 'e' END", num_sql, txt_sql


def render_batch(steps: list[Step], prev: str, own_only: bool = False) -> str:
    """One batch as a SELECT: the new row of every cell this batch computes, after every row of `prev` unless own_only."""
    ctes = ",\n".join(f"{s.cte} AS MATERIALIZED (SELECT {s.expr.sql} AS v, CAST({s.expr.err} AS VARCHAR) AS e, "
                      f"CAST({s.expr.txt} AS VARCHAR) AS t, CAST({s.expr.kind} AS VARCHAR) AS k)" for s in steps)
    rows = []
    for s in steps:
        sheet, row, col = s.key
        typ = s.expr.type
        v, e, t, k = (f"(SELECT {c} FROM {s.cte})" for c in "vetk")
        kind, num_sql, txt_sql = cell_row(typ, v, e, t, k)
        rows.append(f"SELECT {lit_str(sheet).sql}, {row}, {col}, {kind}, {num_sql}, {txt_sql}")
    head = "" if own_only else f"SELECT * FROM {CELLS}\nUNION ALL\n"
    sql = f"WITH\n{ctes}\n{head}" + "\nUNION ALL\n".join(rows)
    return sql.replace(CELLS, prev)


def batch_rows(batch, prev: str) -> str:
    """SQL yielding only the new rows of a batch: a list of Steps (cell batch) or a region batch."""
    if hasattr(batch, "render_rows"):
        return batch.render_rows(prev)
    return render_batch(batch, prev, own_only=True)


def input_rows(book: Book) -> list[tuple]:
    """Every non-formula cell as a row of the cells table (see INPUT_COLUMNS)."""
    rows = []
    for (sheet, row, col), cell in sorted(book.cells.items()):
        if cell.formula is not None or cell.kind == "blank":
            continue
        if cell.kind == "b":
            num_v, txt_v = (1.0 if cell.value else 0.0), None
        elif cell.kind == "n":
            num_v, txt_v = cell.value, None
        else:
            num_v, txt_v = None, cell.value
        rows.append((sheet, row, col, cell.kind, num_v, txt_v))
    return rows


def write_seed(book: Book, path: Path) -> None:
    """The input cells as CSV: the dbt seed, and what the verifier loads (0.1 s for 200k rows)."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(SEED_HEADER)
        for row in input_rows(book):
            w.writerow(["" if x is None else x for x in row])
