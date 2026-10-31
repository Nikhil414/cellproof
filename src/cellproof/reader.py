"""Load a workbook twice: once for formulas, once for the values Excel saved."""
from __future__ import annotations

import dataclasses
import html
from bisect import bisect_left, bisect_right
import datetime
import re
import zipfile
from contextlib import contextmanager
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
from openpyxl.worksheet.formula import ArrayFormula

from cellproof.model import Book, Cell, Kind, Unsupported, Value, a1
from cellproof.tables import Table, expand

ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
          "#CALC!", "#SPILL!", "#FIELD!", "#BUSY!", "#CONNECT!", "#BLOCKED!", "#UNKNOWN!"}  # the last ones: Excel 365
SUPPORTED = {".xlsx", ".xlsm"}
# An .xlsx is a zip of XML. A small file that unzips to gigabytes (a "zip bomb") would exhaust memory in openpyxl.
# ponytail: one total-size limit; raise it if real workbooks ever need more
MAX_UNZIPPED = 2 * 1024**3


class UnreadableWorkbook(Exception):
    pass


ESCAPED = re.compile(r"_x([0-9A-Fa-f]{4})_")


def _text(s: str) -> str:
    """The file format writes a control character as _xHHHH_ (CHAR(13) is _x000D_) and a literal "_x" as _x005F_x:
    one left-to-right pass decodes both, as Excel does."""
    return ESCAPED.sub(lambda m: chr(int(m.group(1), 16)), s) if "_x" in s else s


def _raw_strings(xml_source):
    """openpyxl's read_string_table without its partial unescaping (it drops "x005F_" only, which makes an escaped
    control character and the same text typed literally look alike)."""
    from openpyxl.cell.text import Text
    from openpyxl.xml.constants import SHEET_MAIN_NS
    from openpyxl.xml.functions import iterparse
    strings, tag = [], "{%s}si" % SHEET_MAIN_NS
    for _, node in iterparse(xml_source):
        if node.tag == tag:
            strings.append(Text.from_tree(node).content)
            node.clear()
    return strings


@contextmanager
def _as_saved():
    """Load cells as the file stores them: a date-formatted number stays its serial number (openpyxl's dates lose
    Excel's 29 February 1900, serial 60, and turn a date-formatted -1 into #VALUE!), and text stays escaped for
    _text. Patches openpyxl only while our own load runs."""
    import openpyxl.reader.excel as excel
    import openpyxl.worksheet._reader as sheet_reader
    saved = excel.read_string_table, sheet_reader.from_excel
    excel.read_string_table = _raw_strings
    sheet_reader.from_excel = lambda value, *args, **kwargs: value
    try:
        yield
    finally:
        excel.read_string_table, sheet_reader.from_excel = saved


def _saved(cell, epoch) -> tuple[Value, Kind]:
    v = cell.value
    if v is None:
        if cell.data_type == "str":  # Excel's cached empty string is None in openpyxl.
            return "", "s"
        return None, "blank"
    if cell.data_type == "e":
        return str(v), "e"
    if isinstance(v, bool):  # before int: bool is an int subclass
        return v, "b"
    if isinstance(v, (int, float)):
        return float(v), "n"
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):  # ISO dates (t="d") from other writers
        return float(to_excel(v, epoch)), "n"  # handles Excel's 1900 leap-year bug and the 1904 system
    return _text(str(v)), "s"


def _app(z: zipfile.ZipFile) -> str:
    """The program that last saved the file, as docProps/app.xml names it (printable ASCII, at most 80 characters)."""
    try:
        info = z.getinfo("docProps/app.xml")
    except KeyError:
        return ""
    if info.file_size > 1024**2:
        return ""
    m = re.search(rb"<(?:\w+:)?Application>([^<]*)<", z.read(info))
    return re.sub(r"[^ -~]", "", m.group(1).decode("utf-8", "replace"))[:80].strip() if m else ""


def load(path: str | Path) -> Book:
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED:
        raise UnreadableWorkbook(f"{path.suffix or 'these'} files are not supported in v1 - save the workbook as .xlsx")
    try:
        with zipfile.ZipFile(path) as z:
            size = sum(i.file_size for i in z.infolist())
            app = _app(z)
    except (OSError, zipfile.BadZipFile) as err:
        raise UnreadableWorkbook(f"cannot open {path.name}: {err}") from err
    if size > MAX_UNZIPPED:
        raise UnreadableWorkbook(f"cannot open {path.name}: it unzips to {size / 1024**3:.1f} GB, more than the "
                                 f"{MAX_UNZIPPED // 1024**3} GB limit")
    try:
        with _as_saved():
            fwb = openpyxl.load_workbook(path)
            vwb = openpyxl.load_workbook(path, data_only=True)
    except Exception as err:  # openpyxl raises many types for corrupt or encrypted files
        raise UnreadableWorkbook(f"cannot open {path.name}: {err}") from err

    names = {}
    for name, dn in fwb.defined_names.items():
        try:
            dests = list(dn.destinations)
        except Exception:  # noqa: BLE001 - openpyxl crashes on a name that is a formula of names ("h+i")
            dests = []  # not a plain range: formulas using it are UNRESOLVED ("unknown name")
        if len(dests) == 1:
            sheet, ref = dests[0]
            names[name.upper()] = (sheet, ref.replace("$", ""))

    cells, max_row = {}, {}
    for ws in fwb.worksheets:
        vws = vwb[ws.title]
        max_row[ws.title] = ws.max_row
        # ponytail: openpyxl's private _cells holds only the cells the file has; iter_rows() would build
        # every row x column up to max_row (one styled cell at G1048576 cost 6 s; at Z, 3.8 GB)
        for (r, col), c in sorted(ws._cells.items()):
            if c.value is None:
                continue
            value, kind = _saved(vws.cell(row=r, column=col), fwb.epoch)
            formula, array, ref = None, False, None
            if isinstance(c.value, ArrayFormula):
                formula, array, ref = (c.value.text or "").removeprefix("="), True, c.value.ref
            elif c.data_type == "f":
                formula = c.value.removeprefix("=")
            cells[(ws.title, c.row, c.column)] = Cell(ws.title, c.row, c.column, formula, value, kind, array, ref)
    tables = {}
    for ws in fwb.worksheets:
        for t in ws.tables.values():
            c1, r1, c2, r2 = openpyxl.utils.range_boundaries(t.ref)
            header = 1 if t.headerRowCount is None else int(t.headerRowCount)
            tables[(t.displayName or t.name).upper()] = Table(
                t.displayName or t.name, ws.title, r1, c1, r2, c2, header, int(t.totalsRowCount or 0),
                tuple(col.name for col in t.tableColumns))
    book = Book(cells, [ws.title for ws in fwb.worksheets], names, max_row, fwb.epoch == CALENDAR_MAC_1904, app,
                tables=tables)
    _mark_dynamic(path, book)
    spread_blocks(book)
    rewrite_anchors(book)
    rewrite_tables(book)
    return book


CELL_TAG = re.compile(rb"<c\s[^>]*>")
ATTR = re.compile(rb'\s(r|cm)="([^"]*)"')


def _mark_dynamic(path: Path, book: Book) -> None:
    """Excel marks a dynamic array's anchor cell with cm="..." (cell metadata), which openpyxl does not keep."""
    anchors = {k for k, c in book.cells.items() if c.array}
    if not anchors:
        return
    sheets = {k[0] for k in anchors}
    with zipfile.ZipFile(path) as z:
        wb = z.read("xl/workbook.xml").decode("utf-8", "replace")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8", "replace")
        target = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship\b[^>]*?Id="([^"]+)"[^>]*?Target="([^"]+)"', rels)}
        target.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship\b[^>]*?Target="([^"]+)"[^>]*?Id="([^"]+)"', rels)})
        for m in re.finditer(r'<sheet\b([^>]*)/?>', wb):
            name = re.search(r'\bname="([^"]*)"', m.group(1))
            rid = re.search(r'\br:id="([^"]*)"', m.group(1))
            if not (name and rid) or html.unescape(name.group(1)) not in sheets or rid.group(1) not in target:
                continue
            sheet, part = html.unescape(name.group(1)), target[rid.group(1)].lstrip("/")
            part = part if part.startswith("xl/") else "xl/" + part
            try:
                xml = z.read(part)
            except KeyError:
                continue
            for tag in CELL_TAG.finditer(xml):
                attrs = dict(ATTR.findall(tag.group(0)))
                if b"cm" in attrs and b"r" in attrs:
                    r, c = openpyxl.utils.cell.coordinate_to_tuple(attrs[b"r"].decode())
                    cell = book.cells.get((sheet, r, c))
                    if cell is not None and cell.array:
                        book.cells[cell.key] = dataclasses.replace(cell, dynamic=True)


def spread_blocks(book: Book) -> None:
    """An array formula over a block (A1:C3) is written once, on its top-left cell; the other cells of the block hold
    only Excel's values. Each becomes a formula cell of its own (same text, same block), which the compiler computes
    as this cell's element of the array.
    Only cells the file holds are spread (Excel saves a value in each; a crafted ref as big as the sheet costs
    nothing), and a block over another cell's own formula, or not starting at its anchor, is not read at all."""
    cols: dict[str, list[int]] = {}
    for sheet, col in book._rows_by_col:
        cols.setdefault(sheet, []).append(col)
    for v in cols.values():
        v.sort()
    for key, cell in list(book.cells.items()):
        if not (cell.array and cell.ref and ":" in cell.ref):
            continue
        c1, r1, c2, r2 = openpyxl.utils.range_boundaries(cell.ref)
        have = cols.get(cell.sheet, [])
        inside = []
        for c in have[bisect_left(have, c1):bisect_right(have, c2)]:
            rows = book._rows_by_col[(cell.sheet, c)]
            inside += [(cell.sheet, r, c) for r in rows[bisect_left(rows, r1):bisect_right(rows, r2)]
                       if (r, c) != (cell.row, cell.col)]
        clash = [k for k in inside if book.cells[k].formula is not None]
        if (cell.row, cell.col) != (r1, c1) or clash:
            why = (f"array formula block {cell.ref} overlaps the formula in {a1(clash[0])}" if clash
                   else f"array formula block {cell.ref} does not start at its own cell")
            book.unread[key] = why
            continue
        for k in inside:
            old = book.cells[k]
            book.cells[k] = Cell(cell.sheet, k[1], k[2], cell.formula, old.value, old.kind, True, cell.ref, cell.dynamic)


ANCHOR = re.compile(r"_xlfn\.ANCHORARRAY\(\s*((?:'(?:[^']|'')+'|[A-Za-z0-9_.]+)!)?\$?([A-Za-z]{1,3})\$?([0-9]+)\s*\)")


def rewrite_anchors(book: Book) -> None:
    """A1# (stored _xlfn.ANCHORARRAY(A1)) -> the block Excel spilled A1 into, as the file records it."""
    for key, cell in list(book.cells.items()):
        if cell.formula is None or "ANCHORARRAY" not in cell.formula.upper():
            continue

        def block(m):
            prefix = m.group(1) or ""
            sheet = cell.sheet if not prefix else prefix[:-1].strip("'").replace("''", "'")
            anchor = book.cells.get((sheet, int(m.group(3)), openpyxl.utils.column_index_from_string(m.group(2))))
            if anchor is None or not anchor.array or anchor.row != _top(anchor) or anchor.col != _left(anchor):
                raise Unsupported(f"{m.group(0)}: {m.group(2)}{m.group(3)} is not a spilled array")
            c1, r1, c2, r2 = openpyxl.utils.range_boundaries(anchor.ref or f"{m.group(2)}{m.group(3)}")
            first, last = f"${get_column_letter(c1)}${r1}", f"${get_column_letter(c2)}${r2}"
            return prefix + (first if first == last else f"{first}:{last}")
        try:
            book.cells[key] = dataclasses.replace(cell, formula=ANCHOR.sub(block, cell.formula))
        except Unsupported as err:
            book.unread[key] = str(err)


def _top(cell: Cell) -> int:
    return openpyxl.utils.range_boundaries(cell.ref)[1] if cell.ref else cell.row


def _left(cell: Cell) -> int:
    return openpyxl.utils.range_boundaries(cell.ref)[0] if cell.ref else cell.col


def rewrite_tables(book: Book) -> None:
    """Structured references (Sales[Amount]) -> A1 text in every formula; the ones that cannot be are recorded."""
    for key, cell in list(book.cells.items()):
        if cell.formula is None or "[" not in cell.formula:
            continue
        try:
            book.cells[key] = dataclasses.replace(cell, formula=expand(cell.formula, book.tables, key))
        except Unsupported as err:
            book.unread[key] = str(err)
        except Exception:  # noqa: BLE001 - openpyxl's tokenizer on text it cannot read: the parser says why later
            pass
