"""Compile a Region into one set-based statement.

A region becomes a wide model (a short pivot of the columns it reads, then one CTE per formula column holding the
readable expression) and a mechanical unpivot of that model into long `cells` rows, which is what later batches read
and what the verifier compares. The expressions come from the same Compiler as cell mode, with a different reader for
same-row references, so every Excel rule carries over unchanged.
"""
from __future__ import annotations

import re
from functools import lru_cache

import duckdb
from openpyxl.utils.cell import get_column_letter

from cellproof.compiler import DEEP, INPUT_COLUMNS, Compiler, Step, a1, cell_row
from cellproof.expr import CELLS, NO_ERR, Expr, any_of, errs, finalize, lit_str, num
from cellproof.functions import note_text_numbers, range_sum
from cellproof.model import Book, Unsupported
from cellproof.r1c1 import ROWREF, RUNREF
from cellproof.refs import Area
from cellproof.regions import Region, column_order

# every row from the region's first row to this one (rows are unique and contiguous)
UPTO = "OVER (ORDER BY _.row_num ROWS BETWEEN UNBOUNDED PRECEDING AND CURRENT ROW)"


def _earliest(err: str, pick: str = "arg_min") -> str:
    """The error of the first (arg_min) or last (arg_max) row up to this one that has one, else NULL."""
    return f"{pick}({err}, CASE WHEN {err} IS NOT NULL THEN _.row_num END) {UPTO}"


EPS = 2.0 ** -53
ORDER_REASON = ("a running total whose numbers cancel out (adding and subtracting back to near zero): SQL adds them "
                "in another order than Excel, which can change the last digits")


def _order_safe(steps) -> bool:
    """Whether adding these terms in any order stays within the verifier's tolerance of Excel's left-to-right order.

    steps: (running total so far, sum of absolute values so far) per row. Two summation orders of n terms differ by at
    most about 2 n eps sum|x|; the verifier accepts 1e-9 relative (1e-12 absolute). With terms of one sign sum|x| is
    |total| and the test always holds; a balance that cancels back to zero fails it and stays in cell mode, which adds
    in Excel's order."""
    n = 0
    for total, absolute in steps:
        n += 1
        if 4 * n * EPS * absolute > max(1e-9 * abs(total), 1e-12):
            return False
    return True


def _saved_number(cell):
    """Excel's saved number for a cell, 0 for empty, None when it is an error, text or missing (stop checking there:
    from an error down a running total is an error, so its value no longer matters)."""
    if cell is None or cell.kind == "blank" and cell.formula is None:
        return 0.0
    if cell.kind == "n":
        return cell.value
    if cell.kind == "b":
        return 1.0 if cell.value else 0.0
    return None


def growing_safe(book, region: Region, src: int, start: int) -> bool:
    def steps():
        total = absolute = 0.0
        for r in range(start, region.r2 + 1):
            cell = book.cells.get((region.sheet, r, src))
            if cell is not None and cell.kind == "e":
                return
            if cell is not None and cell.kind == "n":  # SUM skips text, TRUE/FALSE and empty cells
                total += cell.value
                absolute += abs(cell.value)
            if r >= region.r1:
                yield total, absolute
    return _order_safe(steps())


def recurrence_safe(book, region: Region, col: int) -> bool:
    """The terms are read back from Excel's saved totals: term r = X{r} - X{r-1}."""
    def steps():
        prev = _saved_number(book.cells.get((region.sheet, region.r1 - 1, col)))
        if prev is None:
            return
        absolute = abs(prev)
        for r in range(region.r1, region.r2 + 1):
            x = _saved_number(book.cells.get((region.sheet, r, col)))
            if x is None:
                return
            absolute += abs(x - prev)
            prev = x
            yield x, 2 * absolute  # the terms are estimates from rounded totals: count them twice
    return _order_safe(steps())


def window_sum(comp: Compiler, region: Region, src: int, start: int, s: Expr, name: str, layer: list) -> Expr:
    """SUM($B$start:B{r}) as a running window over the source column s; adds its columns to `layer`.

    Like SUM in cell mode: only numbers count (text, TRUE/FALSE skipped), the first error by row wins, and rows above
    the region are one constant read with the cell-mode SUM."""
    if s.type == "num":
        val = s.sql if s.err == NO_ERR else f"CASE WHEN {s.err} IS NULL THEN {s.sql} END"
    elif s.type == "any":
        val = f"CASE WHEN {s.kind} = 'n'{'' if s.err == NO_ERR else f' AND {s.err} IS NULL'} THEN {s.sql} END"
    else:  # text, TRUE/FALSE or errors only: nothing is added
        val = "CAST(NULL AS DOUBLE)"
    layer.append(f"COALESCE(SUM({val}) {UPTO}, CAST(0 AS DOUBLE)) AS {name}")
    err = NO_ERR
    if s.err != NO_ERR:
        layer.append(f"{_earliest(s.err)} AS {name}e")
        err = f"_.{name}e"
    if src not in region.cols:  # an input column: Excel skips numbers stored as text, tell the user once
        note_text_numbers(comp, comp.ready(Area(region.sheet, region.r1, src, region.r2, src)))
    if start < region.r1:
        above = range_sum(comp, Area(region.sheet, start, src, region.r1 - 1, src))
        return Expr(f"({above.sql} + _.{name})", "num", errs(above.err, err))
    return Expr(f"_.{name}", "num", err)


def recurrence(comp: Compiler, region: Region, col: int, e: Expr, name: str, layer: list) -> Expr:
    """X{r-1} + E, X{r-1} - E or E + X{r-1}: the seed above the run plus a running SUM of E (converted like `+`).

    Errors follow Excel's left-to-right rule row by row: with X{r-1} on the left the earliest error wins (the seed's
    first), with it on the right the latest error of E wins (the seed's only when E never failed)."""
    side, op = region.recur[col]
    x = num(e)
    layer.append(f"{x.sql} AS {name}__x")
    werr = NO_ERR
    if x.err != NO_ERR:
        layer.append(f"CAST({x.err} AS VARCHAR) AS {name}__xe")
        werr = _earliest(f"_.{name}__xe", "arg_min" if side == "left" else "arg_max")
    seed = num(comp.cell_value((region.sheet, region.r1 - 1, col)))
    total = f"COALESCE(SUM(_.{name}__x) {UPTO}, CAST(0 AS DOUBLE))"
    return Expr(f"({seed.sql} {op} {total})", "num", errs(seed.err, werr) if side == "left" else errs(werr, seed.err))

RESERVED = {"row_num", "sheet", "src", "wide", "t", "num", "txt", "kind", "col_num", "cells", "all_cells"}


@lru_cache(maxsize=1)
def _keywords() -> frozenset:
    """DuckDB's own list of words that cannot be bare identifiers."""
    rows = duckdb.connect().execute(
        "SELECT keyword_name FROM duckdb_keywords() WHERE keyword_category <> 'unreserved'").fetchall()
    return frozenset(r[0].lower() for r in rows)


WIDE = "\x00wide\x00"  # stands for the relation the unpivot reads: the wide CTE or the wide dbt model


def names_for(book: Book, region: Region, cols: list[int]) -> dict[int, str]:
    """SQL column names: the unique text above the region, else col_<letter>."""
    raw = {}
    for c in cols:
        cell = book.cells.get((region.sheet, region.r1 - 1, c))
        if cell is not None and cell.formula is None and cell.kind == "s":
            raw[c] = re.sub(r"[^0-9a-zA-Z]+", "_", str(cell.value)).strip("_").lower()
    fallback = {c: f"col_{get_column_letter(c).lower()}" for c in cols}
    reserved = RESERVED | _keywords()
    out = {}
    for c in cols:
        name = raw.get(c, "")
        dup = sum(1 for x in raw.values() if x == name) > 1
        bad = (not name or name[0].isdigit() or dup or name in reserved or re.fullmatch(r"col_[a-z]{1,3}", name))
        out[c] = fallback[c] if bad else name
    # last guard: any remaining clash (case-insensitive) sends every colliding header-named column to its fallback
    seen: dict[str, list[int]] = {}
    for c, n in out.items():
        seen.setdefault(n.lower(), []).append(c)
    for group in seen.values():
        if len(group) > 1:
            for c in group:
                out[c] = fallback[c]
    return out


def source_info(comp: Compiler, region: Region, col: int) -> tuple[str, bool]:
    """(type, may_err) of a column the region reads but does not compute; Unsupported if it is not plain."""
    types, may = set(), False
    for r in range(region.r1, region.r2 + 1):
        key = (region.sheet, r, col)
        cell = comp.book.cells.get(key)
        if cell is None:
            raise Unsupported(f"column {get_column_letter(col)} has an empty cell in the region's rows")
        if cell.formula is None:
            if cell.kind not in ("n", "s", "b"):
                raise Unsupported(f"column {get_column_letter(col)} holds {cell.kind} cells")
            types.add({"n": "num", "s": "txt", "b": "bool"}[cell.kind])
        else:
            if key not in comp.done:
                raise Unsupported(f"depends on unresolved cell {a1(key)}")
            typ, err = comp.done[key]
            types.add(typ)
            may = may or err
    if len(types) != 1 or next(iter(types)) not in ("num", "txt", "bool"):
        raise Unsupported(f"column {get_column_letter(col)} mixes value kinds")
    return types.pop(), may


class RegionBatch(list):
    """The Step of every cell of a region, plus the SQL that computes them all at once."""

    def __init__(self, steps, model: str, wide_sql: str, rows: list[str]):
        super().__init__(steps)
        self.model, self.wide_sql, self.rows = model, wide_sql, rows

    def unpivot(self, source: str) -> str:
        return "\nUNION ALL\n".join(row.replace(WIDE, source) for row in self.rows)

    def render_rows(self, prev: str) -> str:
        """Only the new long rows, reading the cells table `prev` (what the verifier appends)."""
        return (f"WITH wide AS (\n{self.wide_sql}\n)\n" + self.unpivot("wide")).replace(CELLS, prev)


ROWREF_REASON = "reads a cell reference directly (functions on single cells are not row-wise yet)"


def model_name(region: Region) -> str:
    slug = re.sub(r"[^0-9a-zA-Z]+", "_", region.sheet).strip("_").lower() or "sheet"
    return f"region_{slug}_{region.r1}_{region.r2}"


def compile_region(comp: Compiler, region: Region):
    """-> (RegionBatch or None, {column: reason it stays in cell mode}). The compiler must have no open batch."""
    assert not comp.open, "compile_region needs an empty open batch"
    book = comp.book
    letter = get_column_letter
    sources = sorted({s for c in region.cols for s in region.reads[c]} - set(region.cols))
    names = names_for(book, region, sorted(set(region.cols) | set(sources)))
    bad_source, src_types = {}, {}
    for s in sources:
        try:
            src_types[s] = source_info(comp, region, s)
        except Unsupported as err:
            bad_source[s] = str(err)
    q = "_."  # every row-wise reference is qualified, so no header name or library alias can shadow it
    rowrefs = {ROWREF.format(s): Expr(q + names[s], t, f"{q}{names[s]}__err" if may else NO_ERR)
               for s, (t, may) in src_types.items()}
    rejected: dict[int, str] = {}
    extras, out_exprs, order = {}, {}, []
    comp.rowrefs = rowrefs
    try:
        for c in column_order(region.cols, region.reads)[0]:
            blocked = sorted(region.reads[c] & (set(rejected) | set(bad_source)))
            if blocked:
                b = blocked[0]
                rejected[c] = bad_source.get(b) or f"reads column {letter(b)}, which stays in cell mode"
                continue
            n = names[c]
            sums: list[str] = []  # the growing SUMs' window columns, then the running total's per-row term
            terms: list[str] = []
            try:
                if not (all(growing_safe(book, region, s, k) for s, k in region.runs.get(c, []))
                        and (c not in region.recur or recurrence_safe(book, region, c))):
                    raise Unsupported(ORDER_REASON)
                for i, (src, start) in enumerate(region.runs.get(c, []), 1):
                    rowrefs[RUNREF.format(src, start)] = window_sum(comp, region, src, start,
                                                                    rowrefs[ROWREF.format(src)], f"{n}__s{i}", sums)
                e = finalize(comp.expr(region.template[c], region.sheet))
                if c in region.recur:
                    e = recurrence(comp, region, c, e, n, terms)
            except Unsupported as err:
                rejected[c] = ROWREF_REASON if "__ROWREF_" in str(err) else str(err)
                continue
            except RecursionError:
                rejected[c] = DEEP
                continue
            extra = [f"{e.sql} AS {n}"]
            if e.err != NO_ERR:
                extra.append(f"CAST({e.err} AS VARCHAR) AS {n}__err")
            if e.type == "any":
                extra += [f"CAST({e.txt} AS VARCHAR) AS {n}__t", f"CAST({e.kind} AS VARCHAR) AS {n}__k"]
            extras[c] = [layer for layer in (sums, terms) if layer] + [extra]
            out_exprs[c] = e
            order.append(c)
            err = f"{q}{n}__err" if e.err != NO_ERR else NO_ERR
            if e.type == "any":
                rowrefs[ROWREF.format(c)] = any_of(f"{q}{n}__k", q + n, f"{q}{n}__t", err)
            elif e.type == "err":
                rowrefs[ROWREF.format(c)] = Expr("NULL", "err", err)
            else:
                rowrefs[ROWREF.format(c)] = Expr(q + n, e.type, err)
    finally:
        comp.rowrefs = None

    sheet_lit = lit_str(region.sheet).sql

    def assemble(accepted: list[int]):
        """(wide_sql, rows, steps) for the accepted columns, which are in compile (dependency) order."""
        used = sorted({s for c in accepted for s in region.reads[c] if s in src_types})
        if used:
            pivot = []
            for s in used:
                typ, may = src_types[s]
                agg = {"num": "MAX(num)", "txt": "MAX(COALESCE(txt, ''))", "bool": "MAX(num)"}[typ]
                filt = f"{agg} FILTER (WHERE col_num = {s})"
                pivot.append(f"({filt} <> 0) AS {names[s]}" if typ == "bool" else f"{filt} AS {names[s]}")
                if may:
                    pivot.append(f"MAX(CASE WHEN kind = 'e' THEN txt END) FILTER (WHERE col_num = {s}) "
                                 f"AS {names[s]}__err")
            src = ("src AS (\n  SELECT row_num,\n    " + ",\n    ".join(pivot) +
                   f"\n  FROM {CELLS} WHERE sheet = {sheet_lit} AND row_num BETWEEN {region.r1} AND {region.r2} "
                   f"AND col_num IN ({', '.join(str(s) for s in used)}) GROUP BY row_num)")
        else:
            src = f"src AS (SELECT row_num FROM range({region.r1}, {region.r2 + 1}) AS t(row_num))"
        layers, prev = [], "src"
        for c in accepted:
            for i, items in enumerate(extras[c], 1):
                cte = names[c] if i == len(extras[c]) else f"{names[c]}__p{i}"  # no SQL name holds "__" otherwise
                layers.append(f"{cte} AS (\n  SELECT *, " + ", ".join(items) + f"\n  FROM {prev} AS _)")
                prev = cte
        cols = [c for c in region.cols if c in accepted]
        select = ["row_num"]
        rows = []
        for c in cols:
            e, n = out_exprs[c], names[c]
            select.append(n)
            if e.err != NO_ERR:
                select.append(f"{n}__err")
            if e.type == "any":
                select += [f"{n}__t", f"{n}__k"]
            v = n if e.type != "err" else "CAST(NULL AS DOUBLE)"
            kind, num_sql, txt_sql = cell_row(e.type, v, f"{n}__err" if e.err != NO_ERR else NO_ERR,
                                              f"{n}__t" if e.type == "any" else "NULL",
                                              f"{n}__k" if e.type == "any" else "NULL")
            rows.append(f"SELECT {sheet_lit}, row_num, {c}, {kind}, {num_sql}, {txt_sql} FROM {WIDE}")
        wide = "WITH\n" + ",\n".join([src] + layers) + f"\nSELECT {', '.join(select)}\nFROM {prev}"
        steps = [Step((region.sheet, r, c), out_exprs[c], "") for c in cols for r in range(region.r1, region.r2 + 1)]
        return wide, rows, steps

    if not order:
        return None, rejected
    accepted = list(order)
    problem = _validate(*assemble(accepted)[:2])
    if problem:  # find the columns DuckDB cannot plan row-wise: add one at a time, keep what plans
        accepted = []
        for c in order:
            blocked = sorted(region.reads[c] & set(rejected))
            if blocked:
                rejected[c] = f"reads column {letter(blocked[0])}, which stays in cell mode"
                continue
            bad = _validate(*assemble(accepted + [c])[:2])
            if bad:
                rejected[c] = f"DuckDB cannot plan this column row-wise ({bad})"
            else:
                accepted.append(c)
    if not accepted:
        return None, rejected
    wide, rows, steps = assemble(accepted)
    return RegionBatch(steps, model_name(region), wide, rows), rejected


_CHECK = None


def _validate(wide_sql: str, rows: list[str]):
    """Plan (EXPLAIN, no data) the region's SQL in both forms it runs in; the first line of the error, or None.

    Form 1 is the verifier's (the cells table); form 2 is dbt's (the cells behind a CTE named all_cells). DuckDB cannot
    plan some correlated subqueries (the lookups); such a column then stays in cell mode instead of crashing the run.
    """
    global _CHECK
    unpivot = "\nUNION ALL\n".join(r.replace(WIDE, "wide") for r in rows)
    dbt_wide = wide_sql.replace("WITH\n", "WITH\nall_cells AS (\n  SELECT * FROM cells UNION ALL SELECT * FROM cells\n),\n", 1)
    forms = [(f"WITH wide AS (\n{wide_sql}\n)\n{unpivot}", "cells"),  # (statement, what the cells are called in it)
             (f"WITH wide AS (\n{dbt_wide}\n)\n{unpivot}", "all_cells")]
    for sql, cells in forms:
        try:
            if _CHECK is None:
                _CHECK = duckdb.connect()
                _CHECK.execute(f"CREATE TABLE cells ({INPUT_COLUMNS})")
            _CHECK.execute("EXPLAIN " + sql.replace(CELLS, cells))
        except Exception as err:  # noqa: BLE001 - any DuckDB failure means "keep this column in cell mode"
            try:
                _CHECK.execute("SELECT 1")
            except Exception:  # noqa: BLE001 - an internal error can invalidate the connection
                _CHECK = None
            return (str(err).splitlines() or ["unknown error"])[0]
    return None


def region_note(region: Region, batch, rejected: dict[int, str]) -> list[str]:
    """Report lines: what became a model and what stayed in cell mode, with the reasons."""
    letters = lambda cols: ", ".join(get_column_letter(c) for c in cols)  # noqa: E731
    lines = []
    if batch is not None:
        done = [c for c in region.cols if c not in rejected]
        lines.append(f"{region.sheet} rows {region.r1}-{region.r2}, columns {letters(done)} -> model {batch.model}")
    for c in sorted(rejected):
        lines.append(f"{region.sheet} column {get_column_letter(c)} rows {region.r1}-{region.r2} stays in cell mode: "
                     f"{rejected[c]}")
    return lines
