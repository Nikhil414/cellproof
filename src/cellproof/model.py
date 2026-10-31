"""Data types shared by every stage."""
from __future__ import annotations

from bisect import bisect_left, bisect_right
from collections import defaultdict
from dataclasses import dataclass, field
from functools import cached_property
from typing import Literal, Union

from openpyxl.utils.cell import get_column_letter

Key = tuple[str, int, int]  # (sheet, row, col)
Kind = Literal["n", "s", "b", "e", "blank"]
Value = Union[float, str, bool, None]


class Unsupported(Exception):
    """A cell cannot be translated. The message becomes the UNRESOLVED reason shown to the user.

    dep: the unresolved cell that stopped the compile, when the reason is that the cell reads one ('depends on
    unresolved cell'); the report attributes the cell to that cell's reason so the reasons can be ranked."""
    def __init__(self, message: str, dep: "Key | None" = None):
        super().__init__(message)
        self.dep = dep


def a1(key: Key) -> str:
    sheet, row, col = key
    return f"{sheet}!{get_column_letter(col)}{row}"


@dataclass(frozen=True)
class Cell:
    sheet: str
    row: int
    col: int
    formula: str | None  # formula text without the leading "=", None for input cells
    value: Value  # Excel-saved value; numbers as float, dates as Excel serial numbers
    kind: Kind  # kind of the saved value
    array: bool = False  # array (Ctrl+Shift+Enter) formula
    ref: str | None = None  # the array formula's ref ("A1" one cell, "A1:A3" a block), None when not an array
    dynamic: bool = False  # a dynamic array (Excel 365 spill): the block is exactly as big as Excel's result

    @property
    def key(self) -> Key:
        return (self.sheet, self.row, self.col)


@dataclass
class Book:
    cells: dict[Key, Cell]
    sheets: list[str]
    names: dict[str, tuple[str, str]] = field(default_factory=dict)  # NAME (upper) -> (sheet, ref)
    max_row: dict[str, int] = field(default_factory=dict)
    date1904: bool = False  # the workbook counts dates from 1904 (old Mac Excel)
    app: str = ""  # the program that last saved the file (docProps/app.xml), "" when the file does not say
    arrays: set[str] = field(default_factory=set)  # hidden sheets of array constants (cellproof.parser.ARRAY_SHEET)
    tables: dict = field(default_factory=dict)  # TABLE NAME (upper) -> cellproof.tables.Table
    unread: dict = field(default_factory=dict)  # formula cell -> why its structured references were not rewritten

    @cached_property
    def _rows_by_col(self) -> dict[tuple[str, int], list[int]]:
        """(sheet, col) -> sorted rows that hold a cell. Built once; build a new Book if cells change."""
        cols = defaultdict(list)
        for sheet, row, col in self.cells:
            cols[(sheet, col)].append(row)
        return {k: sorted(v) for k, v in cols.items()}

    def present(self, sheet: str, r1: int, c1: int, r2: int, c2: int) -> list[Key]:
        """Keys of the non-empty cells inside a rectangle, row-major. Fast on whole-column ranges."""
        found = []
        for col in range(c1, c2 + 1):
            rows = self._rows_by_col.get((sheet, col), [])
            found.extend((sheet, r, col) for r in rows[bisect_left(rows, r1):bisect_right(rows, r2)])
        return sorted(found)


@dataclass(frozen=True)
class Verdict:
    key: Key
    status: Literal["PASS", "FAIL", "UNRESOLVED"]
    expected: Value = None
    actual: Value = None
    reason: str = ""
